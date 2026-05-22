"""Serviço de exportação para PDF e Excel"""
import os
from datetime import datetime
from config import REPORTS_DIR

class ExportService:
    @staticmethod
    def exportar_excel(licitacoes, nome_arquivo=None):
        """Exporta licitações para Excel"""
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
            
            if not nome_arquivo:
                nome_arquivo = f"licitacoes_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            
            caminho_arquivo = os.path.join(REPORTS_DIR, nome_arquivo)
            
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Licitações"
            
            # Headers
            headers = ["ID", "Número", "Título", "Órgão", "Modalidade", 
                      "Data Abertura", "Data Encerramento", "Valor", "Status"]
            ws.append(headers)
            
            # Formatar headers
            header_fill = PatternFill(start_color="1f4788", end_color="1f4788", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF")
            
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")
            
            # Dados
            for licitacao in licitacoes:
                ws.append([
                    licitacao['id'],
                    licitacao['numero'],
                    licitacao['titulo'],
                    licitacao['orgao'],
                    licitacao['modalidade'],
                    licitacao['data_abertura'],
                    licitacao['data_encerramento'],
                    f"R$ {licitacao['valor_estimado']:,.2f}".replace(",", "X").replace(".", ",").replace("X", "."),
                    licitacao['status']
                ])
            
            # Ajustar largura das colunas
            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = (max_length + 2)
                ws.column_dimensions[column_letter].width = adjusted_width
            
            wb.save(caminho_arquivo)
            return caminho_arquivo
        except ImportError:
            return None
    
    @staticmethod
    def exportar_pdf(licitacoes, nome_arquivo=None):
        """Exporta licitações para PDF"""
        try:
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib import colors
            from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.units import inch
            
            if not nome_arquivo:
                nome_arquivo = f"licitacoes_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
            
            caminho_arquivo = os.path.join(REPORTS_DIR, nome_arquivo)
            
            doc = SimpleDocTemplate(caminho_arquivo, pagesize=landscape(A4))
            elements = []
            
            # Estilo
            styles = getSampleStyleSheet()
            title_style = ParagraphStyle(
                'CustomTitle',
                parent=styles['Heading1'],
                fontSize=24,
                textColor=colors.HexColor('#1f4788'),
                alignment=1,
                spaceAfter=30
            )
            
            # Título
            title = Paragraph("Relatório de Licitações", title_style)
            elements.append(title)
            elements.append(Spacer(1, 0.3*inch))
            
            # Tabela
            data = [["ID", "Número", "Título", "Órgão", "Modalidade", "Status"]]
            
            for lic in licitacoes:
                data.append([
                    str(lic['id']),
                    lic['numero'],
                    lic['titulo'][:30],
                    lic['orgao'],
                    lic['modalidade'],
                    lic['status']
                ])
            
            table = Table(data, colWidths=[0.5*inch, 1*inch, 2*inch, 1.5*inch, 1*inch, 0.8*inch])
            
            table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1f4788')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 10),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                ('GRID', (0, 0), (-1, -1), 1, colors.black),
                ('FONTSIZE', (0, 1), (-1, -1), 8),
            ]))
            
            elements.append(table)
            doc.build(elements)
            return caminho_arquivo
        except ImportError:
            return None
